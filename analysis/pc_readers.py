"""Direct (read-only) workbook readers for the weekly labor scorecard on the PC.

Replaces the text-extract parsers (parse_tracker_text / parse_v2_text /
parse_toast_time). Writes the same CSV files labor_eval.py already reads.

Safety: workbooks are opened read_only + data_only and ALWAYS closed in
finally. Nothing is ever saved. All printed output is ASCII.
"""
import csv
import glob
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

DROPBOX = Path.home() / "Dropbox"
CO_STORES = ["98", "99", "101", "246", "256", "491", "516", "680", "1359"]
SD_STORES = ["530", "532", "1360"]
STORES = CO_STORES + SD_STORES
CO_ROOT = DROPBOX / "Colorado" / "Store Reports"
SD_ROOT = DROPBOX / "GM_San Diego" / "Store Reports"
SALES_CACHE = DROPBOX / "Whirl Labor" / "Weekly Sales" / "_gross_cache_2026.json"
TOAST_ROOT = DROPBOX / "performance tracking" / "Toast Files 2026"

V2_KEYS = ["store", "date", "dow", "store_hrs", "parsed", "act_sales", "mgr_fcst", "rec_total", "wtw",
           "earned_sales", "earned_oc", "earned_admin", "earned_train", "earned_total", "act_hrs", "hrs_var", "ot"]
LG_KEYS = ["store", "date", "dow", "store_hrs", "act_sales", "mgr_fcst", "parsed", "fcst_hours_total",
           "sched", "wtw", "ideal_sales", "ideal_oc", "ideal_excl", "visible_excl", "ideal_total", "act_hrs", "hrs_var"]
DOW = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]


def store_dir(store):
    return (SD_ROOT if store in SD_STORES else CO_ROOT) / store


def v2_path(store):
    return store_dir(store) / f"Store {store} Labor and Sales Forecasting 2026 v2.xlsx"


def legacy_path(store):
    return store_dir(store) / f"Store {store} Labor and Sales Forecasting 2026.xlsm"


def n(x):
    """Cell -> float; blank, text and errors become 0.0."""
    if isinstance(x, bool) or x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    return 0.0


def cell(row, col):
    """1-based column from a values row."""
    return row[col - 1] if col - 1 < len(row) else None


def as_date(x):
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    return None


def read_sheets(path, start, end, row_fn):
    """Walk every P* sheet and call row_fn(values_row) for dated day rows in range."""
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for name in wb.sheetnames:
            if not re.fullmatch(r"P\d+", name):
                continue
            for row in wb[name].iter_rows(min_row=4, max_col=36, values_only=True):
                if cell(row, 1) not in DOW:
                    continue
                d = as_date(cell(row, 2))
                if d is None or not (start <= d <= end):
                    continue
                out.append(row_fn(row, d))
    finally:
        wb.close()
    return out


def v2_row(store):
    def fn(row, d):
        act_sales, act_hrs = n(cell(row, 14)), n(cell(row, 27))
        mgr_var = cell(row, 15)
        mgr = n(cell(row, 13))
        r = {"store": store, "date": d.isoformat(), "dow": cell(row, 1), "store_hrs": n(cell(row, 3)),
             "parsed": True, "act_sales": act_sales, "mgr_fcst": mgr if mgr else "",
             "rec_total": n(cell(row, 19)), "wtw": n(cell(row, 20)),
             "earned_sales": n(cell(row, 22)), "earned_oc": n(cell(row, 23)),
             "earned_admin": n(cell(row, 24)), "earned_train": n(cell(row, 25)),
             "earned_total": n(cell(row, 26)), "act_hrs": act_hrs, "hrs_var": n(cell(row, 28)),
             "ot": n(cell(row, 29))}
        return r
    return fn


def legacy_row(store):
    def fn(row, d):
        total, oc = n(cell(row, 34)), n(cell(row, 30))
        sales_ideal = n(cell(row, 29))
        visible = n(cell(row, 31)) + n(cell(row, 32)) + n(cell(row, 33))
        act = n(cell(row, 35))
        has_actuals = n(cell(row, 19)) > 0 and act > 0
        return {"store": store, "date": d.isoformat(), "dow": cell(row, 1), "store_hrs": n(cell(row, 3)),
                "act_sales": n(cell(row, 19)), "mgr_fcst": n(cell(row, 15)), "parsed": has_actuals,
                "fcst_hours_total": n(cell(row, 25)), "sched": n(cell(row, 26)), "wtw": n(cell(row, 27)),
                "ideal_sales": sales_ideal, "ideal_oc": oc,
                "ideal_excl": max(total - sales_ideal - oc, visible), "visible_excl": visible,
                "ideal_total": total, "act_hrs": act, "hrs_var": total - act}
    return fn


def write_csv(path, keys, rows):
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=keys)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in keys})


def build_daily(work, start, end, log=print):
    """Read all 12 stores' v2 and legacy trackers into work/daily_v2.csv + daily_legacy.csv.
    Returns {store: error text} for any store that could not be read."""
    v2_rows, lg_rows, errors = [], [], {}
    for st in STORES:
        for label, path, fn, sink in (("v2", v2_path(st), v2_row(st), v2_rows),
                                      ("legacy", legacy_path(st), legacy_row(st), lg_rows)):
            try:
                rows = read_sheets(path, start, end, fn)
                seen = {(r["store"], r["date"]) for r in sink}
                sink += [r for r in rows if (r["store"], r["date"]) not in seen]
            except Exception as e:  # keep going; report at the end
                errors[f"{st} {label}"] = f"{type(e).__name__}: {e}"
        log(f"read store {st}")
    write_csv(work / "daily_v2.csv", V2_KEYS, v2_rows)
    write_csv(work / "daily_legacy.csv", LG_KEYS, lg_rows)
    return errors


def latest_complete_sunday(v2_csv, stores, today):
    valid = {}
    for r in csv.DictReader(open(v2_csv)):
        valid[(r["store"], r["date"])] = r["parsed"] == "True" and float(r["act_sales"] or 0) > 0 and float(r["act_hrs"] or 0) > 0
    sunday = today - timedelta(days=(today.weekday() + 1) % 7)
    for _ in range(8):
        week = [sunday - timedelta(days=i) for i in range(7)]
        if all(valid.get((s, d.isoformat())) for s in stores for d in week):
            return sunday
        sunday -= timedelta(days=7)
    return None


def build_toast_punches(out_dir, start, end, log=print):
    """Corrected Toast time entries, highest run per folder, one CSV per business date
    (columns labor_eval.load_toast expects: store_code, payable_hours, overtime_hours, job_title)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.csv"):
        old.unlink()
    days = 0
    d = start
    while d <= end:
        folder = TOAST_ROOT / f"{(d + timedelta(days=1)):%b}{(d + timedelta(days=1)).day}"
        runs = sorted(folder.glob("TimeEntries_all_stores_run*.xlsx"),
                      key=lambda p: int(re.search(r"run(\d+)", p.name).group(1)))
        if runs:
            rows = []
            wb = openpyxl.load_workbook(runs[-1], read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                it = ws.iter_rows(values_only=True)
                head = [str(h) for h in next(it)]
                ix = {h: i for i, h in enumerate(head)}
                for r in it:
                    if r[ix["Location Code"]] is None:
                        continue
                    ind = as_date(r[ix["In Date"]])
                    if ind is not None and ind != d:
                        continue  # belongs to another business date
                    rows.append({"store_code": str(int(r[ix["Location Code"]])),
                                 "job_title": str(r[ix["Job Title"]] or ""),
                                 "payable_hours": n(r[ix["Payable Hours"]]),
                                 "overtime_hours": n(r[ix["Overtime Hours"]])})
            finally:
                wb.close()
            write_csv(out_dir / f"{d.isoformat()}.csv", ["store_code", "job_title", "payable_hours", "overtime_hours"], rows)
            days += 1
        d += timedelta(days=1)
    log(f"toast time entries: {days} days")
    return days


def check_sales_cache(end, weeks=4):
    """Names any store/week missing from the Toast weekly sales cache for the window."""
    wk1 = date(2025, 12, 29)
    raw = json.load(open(SALES_CACHE))
    start = end - timedelta(days=7 * weeks - 1)
    missing = []
    for st in STORES:
        for i in range(weeks):
            wk = (start + timedelta(days=7 * i) - wk1).days // 7 + 1
            if str(wk) not in raw.get(st, {}):
                missing.append(f"{st} week {wk}")
    return missing
